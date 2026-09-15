#!/usr/bin/env python3
"""Generate AI_Safety_Measures.xlsx — a structured catalog of AI safety measures.

Run:  python3 generate_ai_safety_workbook.py
Requires: openpyxl (pip install openpyxl)
"""

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
CATEGORY_FILL = PatternFill("solid", fgColor="D9E2F3")
TITLE_FONT = Font(bold=True, size=16, color="1F3864")
THIN_BORDER = Border(
    left=Side(style="thin", color="BFBFBF"),
    right=Side(style="thin", color="BFBFBF"),
    top=Side(style="thin", color="BFBFBF"),
    bottom=Side(style="thin", color="BFBFBF"),
)
WRAP = Alignment(wrap_text=True, vertical="top")


def style_sheet(ws, widths, n_rows, freeze="A2"):
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=2, max_row=n_rows, max_col=len(widths)):
        for cell in row:
            cell.alignment = WRAP
            cell.border = THIN_BORDER
    ws.freeze_panes = freeze
    ws.auto_filter.ref = f"A1:{get_column_letter(len(widths))}{n_rows}"


# ---------------------------------------------------------------------------
# Sheet 2: Safety Measures catalog
# ---------------------------------------------------------------------------
MEASURE_COLS = [
    "ID", "Category", "Safety Measure", "Description",
    "AI Lifecycle Stage", "Risks Addressed", "Implementation Examples",
    "Related Frameworks / Standards", "Status", "Owner", "Notes / Evidence",
]

MEASURES = [
    # --- Governance & Policy ---
    ("GOV-01", "Governance & Policy", "AI governance board / ethics committee",
     "Cross-functional body with authority to review, approve, or halt AI systems based on risk assessments.",
     "All stages", "Loss of oversight; unmanaged risk",
     "Executive AI risk committee; independent ethics advisory board with veto rights",
     "NIST AI RMF (Govern); ISO/IEC 42001"),
    ("GOV-02", "Governance & Policy", "Responsible AI policy & principles",
     "Published organizational principles covering fairness, transparency, accountability, privacy, and safety, with binding internal policies.",
     "All stages", "Inconsistent practices; reputational harm",
     "Public responsible-AI principles; internal policy with mandatory training",
     "OECD AI Principles; ISO/IEC 42001"),
    ("GOV-03", "Governance & Policy", "AI risk management framework adoption",
     "Formal, documented process to identify, measure, mitigate, and monitor AI risks across the portfolio of AI systems.",
     "All stages", "Unidentified or unmanaged risks",
     "NIST AI RMF implementation; AI risk register reviewed quarterly",
     "NIST AI RMF; ISO/IEC 23894"),
    ("GOV-04", "Governance & Policy", "Frontier safety / responsible scaling policy",
     "Predefined capability thresholds that trigger stronger safeguards, evaluations, or pauses before training or deploying more capable models.",
     "Design; Training; Deployment", "Dangerous emergent capabilities; loss of control",
     "Anthropic Responsible Scaling Policy; OpenAI Preparedness Framework; Google DeepMind Frontier Safety Framework",
     "Frontier AI Safety Commitments (Seoul 2024)"),
    ("GOV-05", "Governance & Policy", "Regulatory compliance mapping",
     "Systematic mapping of AI systems to applicable regulations, including risk-tier classification and conformity assessment.",
     "All stages", "Legal / regulatory penalties",
     "EU AI Act risk classification; sectoral compliance reviews (health, finance)",
     "EU AI Act; sector regulations"),
    ("GOV-06", "Governance & Policy", "Third-party / vendor AI risk assessment",
     "Due diligence on external models, datasets, and AI services before procurement or integration.",
     "Design; Deployment", "Supply-chain risk; inherited vulnerabilities",
     "Vendor questionnaires; model provenance checks; contractual safety clauses",
     "ISO/IEC 42001; NIST AI RMF (Map)"),
    ("GOV-07", "Governance & Policy", "AI incident response plan",
     "Documented procedures for detecting, escalating, containing, and remediating AI failures or misuse, including external notification.",
     "Deployment; Monitoring", "Prolonged harm from incidents",
     "AI-specific runbooks; severity classification; postmortems; regulator notification workflow",
     "NIST AI RMF (Manage); EU AI Act Art. 73"),
    ("GOV-08", "Governance & Policy", "Internal reporting & whistleblower channels",
     "Protected channels for employees to raise safety concerns about AI systems without retaliation.",
     "All stages", "Suppressed safety signals",
     "Anonymous reporting hotline; safety concern escalation policy",
     "ISO/IEC 42001; EU Whistleblower Directive"),

    # --- Data Safety ---
    ("DAT-01", "Data Safety", "Data provenance & lineage tracking",
     "Recording the origin, licensing, and transformation history of all training and evaluation data.",
     "Data collection; Training", "IP infringement; untraceable errors",
     "Dataset registries; datasheets for datasets; C2PA content credentials",
     "EU AI Act Art. 10; Datasheets for Datasets"),
    ("DAT-02", "Data Safety", "PII detection & minimization",
     "Automated scanning and removal or anonymization of personal data in training corpora and logs.",
     "Data collection; Training; Monitoring", "Privacy violations; data leakage",
     "PII scrubbing pipelines; pseudonymization; retention limits",
     "GDPR; ISO/IEC 27701"),
    ("DAT-03", "Data Safety", "Training data bias audit",
     "Statistical analysis of dataset representativeness and label quality across demographic groups and domains.",
     "Data collection; Training", "Discriminatory model behavior",
     "Demographic coverage reports; slice-based label QA",
     "NIST AI RMF (Measure); ISO/IEC TR 24027"),
    ("DAT-04", "Data Safety", "Data poisoning defenses",
     "Filtering, deduplication, and anomaly detection to prevent adversarial or low-quality data from corrupting training.",
     "Data collection; Training", "Backdoored or degraded models",
     "Source allowlists; near-duplicate detection; outlier filtering; trusted-data pipelines",
     "MITRE ATLAS; OWASP ML Top 10"),
    ("DAT-05", "Data Safety", "Differential privacy in training",
     "Mathematical privacy guarantees that limit what a model can memorize about any individual training example.",
     "Training", "Training-data extraction; memorization",
     "DP-SGD; privacy budget accounting; memorization audits",
     "NIST SP 800-226"),

    # --- Model Development & Evaluation ---
    ("DEV-01", "Model Development & Evaluation", "Safety alignment training",
     "Techniques that steer model behavior toward helpful, harmless outputs and refusal of harmful requests.",
     "Training", "Harmful or toxic outputs",
     "RLHF; Constitutional AI; safety-focused fine-tuning; refusal training",
     "Frontier AI Safety Commitments"),
    ("DEV-02", "Model Development & Evaluation", "Red teaming / adversarial testing",
     "Structured attempts by internal or external experts to elicit harmful, biased, or policy-violating behavior before release.",
     "Pre-deployment", "Undiscovered failure modes; jailbreaks",
     "Internal red teams; external expert engagements; automated adversarial prompt generation",
     "NIST AI RMF (Measure); EU AI Act Art. 55; OWASP LLM Top 10"),
    ("DEV-03", "Model Development & Evaluation", "Dangerous capability evaluations",
     "Testing for capabilities that could enable severe harm (cyber offense, biological/chemical weapon uplift, autonomous replication, deception).",
     "Pre-deployment", "Catastrophic misuse; loss of control",
     "CBRN uplift evals; cyber-offense benchmarks; autonomy evals; third-party evaluation (e.g., national AI safety institutes)",
     "Frontier AI Safety Commitments; company frontier safety policies"),
    ("DEV-04", "Model Development & Evaluation", "Bias & fairness evaluation",
     "Quantitative evaluation of model performance and error rates across demographic groups and protected attributes.",
     "Pre-deployment; Monitoring", "Discriminatory outcomes",
     "Disaggregated accuracy metrics; fairness benchmarks (e.g., BBQ); counterfactual testing",
     "NIST SP 1270; ISO/IEC TR 24027"),
    ("DEV-05", "Model Development & Evaluation", "Robustness testing",
     "Evaluating model stability under distribution shift, adversarial perturbations, and edge-case inputs.",
     "Pre-deployment", "Unpredictable failures in the wild",
     "Adversarial example testing; stress tests; out-of-distribution benchmarks",
     "ISO/IEC 24029; MITRE ATLAS"),
    ("DEV-06", "Model Development & Evaluation", "Hallucination / factuality evaluation",
     "Measuring the rate of fabricated or incorrect claims, especially in high-stakes domains.",
     "Pre-deployment; Monitoring", "Misinformation; unsafe advice",
     "Factuality benchmarks; citation-grounding checks; domain-expert review",
     "NIST AI RMF (Measure)"),
    ("DEV-07", "Model Development & Evaluation", "Interpretability & explainability tooling",
     "Methods to understand model internals and explain individual decisions to developers, auditors, and affected users.",
     "Training; Pre-deployment; Deployment", "Opacity; undetected misalignment",
     "Feature attribution (SHAP/LIME); mechanistic interpretability research; decision rationale generation",
     "EU AI Act Art. 13; NIST AI RMF"),
    ("DEV-08", "Model Development & Evaluation", "Safety benchmark regression suite",
     "Automated safety test suites run on every model version to prevent regressions before promotion.",
     "Training; Pre-deployment", "Silent safety regressions",
     "CI-integrated safety evals; release gates tied to benchmark thresholds",
     "Internal MLOps best practice"),

    # --- Deployment Safeguards ---
    ("DEP-01", "Deployment Safeguards", "Input/output guardrails & content filtering",
     "Real-time classifiers and rule systems that block harmful prompts and filter unsafe model outputs.",
     "Deployment", "Harmful content generation; policy violations",
     "Moderation APIs; toxicity classifiers; topic blocklists; safety system prompts",
     "OWASP LLM Top 10; platform usage policies"),
    ("DEP-02", "Deployment Safeguards", "Jailbreak & prompt injection defenses",
     "Detection and mitigation of attempts to bypass safety training via crafted prompts or injected instructions in retrieved content.",
     "Deployment", "Safety bypass; data exfiltration; agent hijacking",
     "Prompt injection classifiers; instruction hierarchy; sandboxed tool use; input sanitization",
     "OWASP LLM Top 10 (LLM01); MITRE ATLAS"),
    ("DEP-03", "Deployment Safeguards", "Human-in-the-loop for high-stakes decisions",
     "Mandatory human review and override capability for consequential decisions (medical, legal, financial, safety-critical).",
     "Deployment", "Automated harm at scale",
     "Approval workflows; confidence-threshold escalation; human override UI",
     "EU AI Act Art. 14; NIST AI RMF (Govern)"),
    ("DEP-04", "Deployment Safeguards", "Staged rollout & canary deployment",
     "Gradual exposure of new models to limited traffic with monitoring before full release.",
     "Deployment", "Widespread impact of undetected flaws",
     "Percentage-based rollouts; A/B safety comparisons; feature flags",
     "MLOps best practice"),
    ("DEP-05", "Deployment Safeguards", "Kill switch & rollback capability",
     "Ability to rapidly disable an AI system or revert to a previous safe version when serious issues are detected.",
     "Deployment; Monitoring", "Prolonged exposure to failures",
     "One-click model rollback; emergency shutdown procedures; circuit breakers",
     "EU AI Act; NIST AI RMF (Manage)"),
    ("DEP-06", "Deployment Safeguards", "Access controls & usage tiering",
     "Restricting powerful capabilities to vetted users, with authentication, authorization, and rate limits.",
     "Deployment", "Misuse by malicious actors",
     "API key vetting; KYC for high-risk capabilities; rate limiting; capability gating",
     "ISO/IEC 27001; frontier safety policies"),
    ("DEP-07", "Deployment Safeguards", "AI output watermarking & provenance",
     "Embedding detectable signals or metadata in AI-generated content to support authenticity verification.",
     "Deployment", "Deepfakes; disinformation",
     "SynthID; C2PA content credentials; statistical text watermarking",
     "EU AI Act Art. 50; C2PA"),
    ("DEP-08", "Deployment Safeguards", "AI use disclosure to users",
     "Clearly informing users when they are interacting with an AI system or consuming AI-generated content.",
     "Deployment", "Deception; loss of trust",
     "Chatbot disclosure banners; AI-generated content labels",
     "EU AI Act Art. 50; FTC guidance"),
    ("DEP-09", "Deployment Safeguards", "Agentic AI containment",
     "Sandboxing, permission scoping, and action approval for AI agents that can take autonomous actions or use tools.",
     "Deployment", "Unintended autonomous actions",
     "Sandboxed execution; tool allowlists; spend limits; human approval for irreversible actions",
     "OWASP LLM Top 10 (Excessive Agency)"),

    # --- Monitoring & Response ---
    ("MON-01", "Monitoring & Response", "Production behavior monitoring",
     "Continuous tracking of model outputs, refusal rates, safety classifier triggers, and quality metrics in production.",
     "Monitoring", "Undetected degradation or misuse",
     "Safety dashboards; automated alerting on anomaly thresholds",
     "NIST AI RMF (Manage); ISO/IEC 42001"),
    ("MON-02", "Monitoring & Response", "Model & data drift detection",
     "Statistical monitoring for shifts in input distributions or model performance that may invalidate safety assumptions.",
     "Monitoring", "Silent performance decay",
     "Drift metrics (PSI, KL divergence); scheduled re-evaluation; retraining triggers",
     "MLOps best practice"),
    ("MON-03", "Monitoring & Response", "Abuse & misuse monitoring",
     "Detecting coordinated misuse, policy-violating usage patterns, and threat-actor activity across the platform.",
     "Monitoring", "Weaponization; fraud; influence operations",
     "Usage pattern analysis; threat intelligence integration; account enforcement",
     "Platform usage policies; frontier safety policies"),
    ("MON-04", "Monitoring & Response", "User feedback & incident reporting channels",
     "Mechanisms for users and affected parties to report harmful outputs or unexpected behavior.",
     "Monitoring", "Missed real-world harm signals",
     "In-product report buttons; feedback triage workflows; appeal processes",
     "EU AI Act; DSA (for platforms)"),
    ("MON-05", "Monitoring & Response", "Vulnerability disclosure & bug bounty",
     "Formal programs inviting external researchers to report model vulnerabilities, jailbreaks, and safety flaws.",
     "Monitoring", "Unreported exploitable weaknesses",
     "Model safety bug bounties; coordinated disclosure policies",
     "ISO/IEC 29147; industry practice"),
    ("MON-06", "Monitoring & Response", "Comprehensive audit logging",
     "Tamper-evident logs of model inputs, outputs, decisions, and administrative actions to support investigation and accountability.",
     "Deployment; Monitoring", "Inability to investigate incidents",
     "Immutable log stores; retention policies; access-controlled review tooling",
     "EU AI Act Art. 12; SOC 2"),

    # --- Security ---
    ("SEC-01", "Security", "Model weight & IP security",
     "Protecting model weights and proprietary training artifacts from theft or unauthorized exfiltration.",
     "Training; Deployment", "Model theft; uncontrolled proliferation",
     "Weight encryption; hardware security modules; insider-threat programs; RAND SL frameworks",
     "RAND securing-weights framework; frontier safety policies"),
    ("SEC-02", "Security", "ML supply chain security",
     "Securing the pipeline of datasets, libraries, pretrained models, and infrastructure against tampering.",
     "All stages", "Backdoored dependencies or models",
     "Model/dataset signing; SBOM for ML; dependency scanning; trusted registries",
     "NIST SSDF; MITRE ATLAS; OWASP ML Top 10"),
    ("SEC-03", "Security", "Adversarial ML attack defenses",
     "Defenses against evasion, extraction, inversion, and membership-inference attacks on deployed models.",
     "Deployment", "Model extraction; privacy attacks",
     "Query anomaly detection; output perturbation; rate limits; distillation defenses",
     "MITRE ATLAS; NIST AI 100-2"),
    ("SEC-04", "Security", "Secure inference infrastructure",
     "Hardening of serving infrastructure, including isolation of customer data and confidential computing where warranted.",
     "Deployment", "Data breaches; cross-tenant leakage",
     "Confidential computing / TEEs; tenant isolation; encryption in transit and at rest",
     "ISO/IEC 27001; SOC 2"),

    # --- Transparency & Accountability ---
    ("TRA-01", "Transparency & Accountability", "Model cards / system cards",
     "Standardized documentation of a model's intended use, capabilities, limitations, evaluation results, and safety mitigations.",
     "Pre-deployment; Deployment", "Misuse from misunderstanding; opacity",
     "Published model cards; system cards accompanying major releases",
     "Model Cards (Mitchell et al.); EU AI Act Annex IV"),
    ("TRA-02", "Transparency & Accountability", "Independent third-party audits",
     "External assessment of AI systems, safety processes, and claims by qualified independent auditors.",
     "Pre-deployment; Monitoring", "Unverified self-assessment",
     "Algorithmic audits; ISO/IEC 42001 certification; national AI safety institute evaluations",
     "ISO/IEC 42001; EU AI Act conformity assessment"),
    ("TRA-03", "Transparency & Accountability", "Safety research publication & sharing",
     "Publishing safety findings, incident learnings, and best practices to raise industry-wide standards.",
     "All stages", "Slow industry-wide safety progress",
     "Safety research papers; incident databases (e.g., AI Incident Database); industry forums",
     "Frontier Model Forum; Partnership on AI"),
    ("TRA-04", "Transparency & Accountability", "Impact assessments",
     "Structured pre-deployment assessment of potential effects on individuals, groups, and society, with documented mitigations.",
     "Design; Pre-deployment", "Unanticipated societal harm",
     "Algorithmic impact assessments; DPIAs; fundamental rights impact assessments",
     "EU AI Act Art. 27; GDPR Art. 35; Canada AIA"),
]

# ---------------------------------------------------------------------------
# Sheet 3: Frameworks & Standards
# ---------------------------------------------------------------------------
FRAMEWORK_COLS = ["Framework / Standard", "Issuing Body", "Year", "Type", "Scope & Purpose"]
FRAMEWORKS = [
    ("NIST AI Risk Management Framework (AI RMF 1.0)", "NIST (US)", "2023", "Voluntary framework",
     "Risk management functions (Govern, Map, Measure, Manage) for trustworthy AI across the lifecycle; Generative AI Profile added 2024."),
    ("EU AI Act (Regulation 2024/1689)", "European Union", "2024", "Binding regulation",
     "Risk-tiered obligations: prohibited practices, high-risk system requirements (data, oversight, logging, robustness), transparency duties, and general-purpose AI model rules."),
    ("ISO/IEC 42001", "ISO/IEC", "2023", "Certifiable standard",
     "AI management system (AIMS) requirements — the organizational governance layer for responsible AI development and use."),
    ("ISO/IEC 23894", "ISO/IEC", "2023", "Guidance standard",
     "AI-specific risk management guidance aligned with ISO 31000."),
    ("OECD AI Principles", "OECD", "2019 (rev. 2024)", "Intergovernmental principles",
     "High-level principles: inclusive growth, human rights, transparency, robustness, and accountability; basis for many national policies."),
    ("Frontier AI Safety Commitments", "Seoul AI Summit signatories", "2024", "Voluntary commitments",
     "Frontier labs commit to publish safety frameworks with capability thresholds, mitigations, and conditions to pause development."),
    ("Company frontier safety policies (RSP / Preparedness / FSF)", "Anthropic, OpenAI, Google DeepMind, et al.", "2023–", "Company policy",
     "Capability-threshold-based safeguards: dangerous capability evals, security standards for weights, and deployment gates."),
    ("NIST AI 100-2 (Adversarial ML taxonomy)", "NIST (US)", "2024", "Technical report",
     "Taxonomy and terminology of adversarial ML attacks and mitigations (evasion, poisoning, extraction, inference)."),
    ("MITRE ATLAS", "MITRE", "2021–", "Threat knowledge base",
     "Adversary tactics and techniques targeting AI systems, modeled after ATT&CK; supports AI threat modeling and red teaming."),
    ("OWASP Top 10 for LLM Applications", "OWASP", "2023 (rev. 2025)", "Community guidance",
     "Top security risks for LLM apps — prompt injection, insecure output handling, data poisoning, excessive agency — with mitigations."),
    ("ISO/IEC 27001 / SOC 2", "ISO/IEC; AICPA", "Ongoing", "Certifiable standards",
     "General information-security management and trust-services controls that underpin AI infrastructure security."),
    ("C2PA Content Credentials", "Coalition for Content Provenance and Authenticity", "2022–", "Technical specification",
     "Cryptographically signed provenance metadata for media, enabling AI-generated content labeling."),
    ("G7 Hiroshima Process Code of Conduct", "G7", "2023", "Voluntary code",
     "International code of conduct for organizations developing advanced AI systems, covering risk identification, incident reporting, and provenance."),
]

# ---------------------------------------------------------------------------
# Sheet 4: Risk Categories
# ---------------------------------------------------------------------------
RISK_COLS = ["Risk Category", "Description", "Example Harms", "Primary Mitigating Measure Categories"]
RISKS = [
    ("Bias & discrimination", "Systematically worse outcomes for certain groups due to skewed data or design choices.",
     "Discriminatory hiring, lending, or policing decisions", "Data Safety; Model Development & Evaluation"),
    ("Privacy violations", "Exposure or inference of personal data through training data memorization or inference attacks.",
     "Training-data extraction; re-identification; surveillance", "Data Safety; Security"),
    ("Misinformation & hallucination", "Confident generation of false or fabricated content.",
     "Unsafe medical/legal advice; fabricated citations; deepfakes", "Model Development & Evaluation; Deployment Safeguards"),
    ("Malicious misuse", "Deliberate use of AI capabilities to cause harm.",
     "Cyberattack automation; bioweapon uplift; fraud; influence operations", "Governance & Policy; Deployment Safeguards; Monitoring & Response"),
    ("Security vulnerabilities", "Attacks on the AI system itself or via its integrations.",
     "Prompt injection; model theft; data poisoning; supply-chain compromise", "Security; Deployment Safeguards"),
    ("Lack of robustness", "Failures under distribution shift, adversarial inputs, or edge cases.",
     "Autonomous vehicle misclassification; brittle medical models", "Model Development & Evaluation; Monitoring & Response"),
    ("Opacity & unaccountability", "Inability to explain, audit, or assign responsibility for AI decisions.",
     "Unappealable automated decisions; unverifiable safety claims", "Transparency & Accountability; Governance & Policy"),
    ("Loss of control / misalignment", "Advanced systems pursuing unintended objectives or resisting oversight.",
     "Deceptive behavior; unsafe autonomous actions at scale", "Governance & Policy; Model Development & Evaluation; Deployment Safeguards"),
    ("Overreliance & automation bias", "Humans deferring to AI beyond its competence.",
     "Skill atrophy; unchecked erroneous decisions in high-stakes settings", "Deployment Safeguards; Transparency & Accountability"),
    ("Societal & economic impacts", "Broader externalities of AI deployment.",
     "Labor displacement; concentration of power; environmental cost", "Governance & Policy; Transparency & Accountability"),
]

STATUS_OPTIONS = ["Not Started", "Planned", "In Progress", "Implemented", "Verified", "N/A"]


def build_overview(wb):
    ws = wb.active
    ws.title = "Overview"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 95

    ws["B2"] = "AI Safety Measures Catalog"
    ws["B2"].font = TITLE_FONT
    ws["B3"] = "A structured inventory of AI safety measures, the risks they mitigate, and the frameworks they map to."
    ws["B3"].font = Font(italic=True, color="595959")

    rows = [
        ("Sheet", "Contents"),
        ("Safety Measures", f"{len(MEASURES)} measures across 7 categories, with descriptions, lifecycle stage, risks addressed, implementation examples, framework mappings, and Status/Owner tracking columns (Status has a dropdown)."),
        ("Frameworks & Standards", f"{len(FRAMEWORKS)} key AI safety and governance frameworks, standards, and regulations."),
        ("Risk Categories", f"{len(RISKS)} AI risk categories with example harms and the measure categories that mitigate them."),
    ]
    start = 5
    for i, (a, b) in enumerate(rows):
        ra, rb = ws.cell(row=start + i, column=2, value=a), ws.cell(row=start + i, column=3, value=b)
        ra.alignment = WRAP
        rb.alignment = WRAP
        ra.border = THIN_BORDER
        rb.border = THIN_BORDER
        if i == 0:
            ra.fill = HEADER_FILL
            rb.fill = HEADER_FILL
            ra.font = HEADER_FONT
            rb.font = HEADER_FONT

    usage_start = start + len(rows) + 2
    ws.cell(row=usage_start, column=2, value="How to use").font = Font(bold=True, size=12, color="1F3864")
    usage = [
        "1. Review the Safety Measures sheet and filter by Category or AI Lifecycle Stage using the header filters.",
        "2. For each measure, set Status (dropdown: " + ", ".join(STATUS_OPTIONS) + ") and assign an Owner.",
        "3. Record supporting evidence or links in the Notes / Evidence column.",
        "4. Use Risk Categories to check coverage: every relevant risk should map to implemented measures.",
        "5. Use Frameworks & Standards to align the catalog with the regulations and standards that apply to you.",
    ]
    for i, line in enumerate(usage, start=1):
        c = ws.cell(row=usage_start + i, column=2, value=line)
        c.alignment = WRAP
        ws.merge_cells(start_row=usage_start + i, start_column=2, end_row=usage_start + i, end_column=3)

    note_row = usage_start + len(usage) + 2
    c = ws.cell(row=note_row, column=2,
                value="Note: This catalog is a general-purpose reference compiled from public frameworks and industry practice (as of 2026). Tailor it to your organization's context, risk appetite, and applicable regulations.")
    c.font = Font(italic=True, color="808080")
    c.alignment = WRAP
    ws.merge_cells(start_row=note_row, start_column=2, end_row=note_row, end_column=3)


def build_measures(wb):
    ws = wb.create_sheet("Safety Measures")
    ws.append(MEASURE_COLS)
    for m in MEASURES:
        ws.append(list(m) + ["Not Started", "", ""])

    n_rows = len(MEASURES) + 1
    style_sheet(ws, widths=[9, 22, 30, 52, 18, 26, 42, 30, 13, 14, 26], n_rows=n_rows, freeze="D2")

    for row in ws.iter_rows(min_row=2, max_row=n_rows, min_col=2, max_col=2):
        for cell in row:
            cell.fill = CATEGORY_FILL

    dv = DataValidation(type="list", formula1='"' + ",".join(STATUS_OPTIONS) + '"', allow_blank=True,
                        showDropDown=False)
    dv.error = "Please select a value from the list."
    dv.errorTitle = "Invalid status"
    ws.add_data_validation(dv)
    dv.add(f"I2:I{n_rows}")


def build_frameworks(wb):
    ws = wb.create_sheet("Frameworks & Standards")
    ws.append(FRAMEWORK_COLS)
    for f in FRAMEWORKS:
        ws.append(list(f))
    style_sheet(ws, widths=[42, 30, 14, 22, 85], n_rows=len(FRAMEWORKS) + 1)


def build_risks(wb):
    ws = wb.create_sheet("Risk Categories")
    ws.append(RISK_COLS)
    for r in RISKS:
        ws.append(list(r))
    style_sheet(ws, widths=[28, 55, 48, 45], n_rows=len(RISKS) + 1)


def main():
    wb = Workbook()
    build_overview(wb)
    build_measures(wb)
    build_frameworks(wb)
    build_risks(wb)
    out = "AI_Safety_Measures.xlsx"
    wb.save(out)
    print(f"Wrote {out}: {len(MEASURES)} measures, {len(FRAMEWORKS)} frameworks, {len(RISKS)} risk categories.")


if __name__ == "__main__":
    main()
