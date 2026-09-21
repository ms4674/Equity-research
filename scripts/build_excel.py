#!/usr/bin/env python3
"""Build data/marketplace_bookings_2025.xlsx from data/marketplace_bookings_2025.csv.

Produces two sheets:
- "Summary": sector subtotals driven by live SUMIFS formulas over the Data sheet,
  plus the ex-bed-bank adjusted totals and methodology notes.
- "Data": the full per-company dataset with filters, frozen header, and
  hyperlinked sources.
"""

import csv
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
CSV_PATH = ROOT / "data" / "marketplace_bookings_2025.csv"
XLSX_PATH = ROOT / "data" / "marketplace_bookings_2025.xlsx"

# Bases treated as company-sourced rather than analyst/third-party estimates.
DISCLOSED_BASES = {"disclosed", "disclosed (stale)", "company-announced", "management-stated approx"}
B2B_BED_BANKS = ["HBX Group (Hotelbeds)", "WEB Travel Group (WebBeds)"]

SECTORS = [
    ("Delivery", "Food & grocery delivery"),
    ("Ridesharing", "Ridesharing / ride-hailing"),
    ("Online travel", "Online travel"),
]

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TOTAL_FILL = PatternFill("solid", fgColor="D9E2F3")
MEMO_FONT = Font(italic=True, color="808080")
THIN_BORDER = Border(bottom=Side(style="thin", color="BFBFBF"))

DATA_HEADERS = [
    ("Sector", 14),
    ("Company", 42),
    ("Ticker", 26),
    ("Metric", 36),
    ("Period", 22),
    ("Value (USD bn)", 14),
    ("Basis", 24),
    ("Basis bucket", 13),
    ("In total", 9),
    ("Notes", 80),
    ("Source", 40),
]


def build_data_sheet(ws, rows):
    for col, (title, width) in enumerate(DATA_HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = width

    for r, row in enumerate(rows, start=2):
        bucket = "Disclosed" if row["basis"] in DISCLOSED_BASES else "Estimate"
        value = float(row["value_usd_bn"]) if row["value_usd_bn"] else None
        values = [
            row["sector"], row["company"], row["ticker"], row["metric"], row["period"],
            value, row["basis"], bucket if value is not None else "", row["in_total"], row["notes"],
        ]
        for col, val in enumerate(values, start=1):
            cell = ws.cell(row=r, column=col, value=val)
            cell.border = THIN_BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=(col in (4, 10)))
            if row["in_total"] == "memo":
                cell.font = MEMO_FONT
        ws.cell(row=r, column=6).number_format = "0.0"

        src = ws.cell(row=r, column=11, value=urlparse(row["source"]).netloc or row["source"])
        src.hyperlink = row["source"]
        src.font = Font(color="0563C1", underline="single", italic=(row["in_total"] == "memo"))
        src.border = THIN_BORDER
        src.alignment = Alignment(vertical="top")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:K{len(rows) + 1}"


def build_summary_sheet(ws):
    ws.column_dimensions["A"].width = 44
    for col in "BCD":
        ws.column_dimensions[col].width = 15

    ws["A1"] = "Global Marketplace Bookings: Delivery, Ridesharing & Online Travel (2025)"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = ("Gross transaction value (GOV / GMV / GTV / gross bookings / GBV), USD billions, "
                "latest completed fiscal year. Compiled September 2026.")
    ws["A2"].font = Font(italic=True, color="595959")

    for col, title in zip("ABCD", ["Sector", "Disclosed", "Estimated", "Subtotal"]):
        cell = ws[f"{col}4"]
        cell.value = title
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    row = 5
    for key, label in SECTORS:
        ws[f"A{row}"] = label
        ws[f"B{row}"] = (f'=SUMIFS(Data!$F:$F,Data!$A:$A,"{key}",'
                         f'Data!$H:$H,"Disclosed",Data!$I:$I,"yes")')
        ws[f"C{row}"] = (f'=SUMIFS(Data!$F:$F,Data!$A:$A,"{key}",'
                         f'Data!$H:$H,"Estimate",Data!$I:$I,"yes")')
        ws[f"D{row}"] = f"=B{row}+C{row}"
        row += 1

    total_row = row
    ws[f"A{total_row}"] = "Grand total"
    for col in "BCD":
        ws[f"{col}{total_row}"] = f"={col}5+{col}6+{col}7"
    for col in "ABCD":
        ws[f"{col}{total_row}"].fill = TOTAL_FILL
        ws[f"{col}{total_row}"].font = Font(bold=True)

    bed_bank_formula = "+".join(
        f'SUMIFS(Data!$F:$F,Data!$B:$B,"{name}")' for name in B2B_BED_BANKS
    )
    adj_row = total_row + 2
    ws[f"A{adj_row}"] = "Less: B2B bed banks (HBX, WebBeds)"
    ws[f"D{adj_row}"] = f"=-({bed_bank_formula})"
    ws[f"A{adj_row + 1}"] = "Consumer-facing total (ex-bed banks)"
    ws[f"D{adj_row + 1}"] = f"=D{total_row}+D{adj_row}"
    ws[f"A{adj_row + 1}"].font = Font(bold=True)
    ws[f"D{adj_row + 1}"].font = Font(bold=True)

    for r in range(5, adj_row + 2):
        for col in "BCD":
            ws[f"{col}{r}"].number_format = "#,##0.0"

    notes = [
        "Notes:",
        "- Subtotals are live SUMIFS over the Data sheet; edit values or the 'In total' flag there to update.",
        "- 'Disclosed' includes company-reported figures plus management-stated (Trip.com) and "
        "company-announced (GetYourGuide) values; 'Estimate' covers analyst/third-party triangulations.",
        "- Memo rows (talabat inside Delivery Hero; Careem, FreeNow, Wolt, Kakao Mobility) are excluded "
        "from totals to avoid double counting.",
        "- China delivery (Meituan, Ele.me/Taobao, JD) does not disclose absolute transaction value; "
        "~$275B of the delivery subtotal is estimated, with a roughly +/-$50B plausible range.",
        "- Stale entries: Tongcheng and Despegar are FY2024; Hopper is a CY2024 third-party estimate.",
        "- Full methodology and caveats: reports/global-marketplace-bookings-2025.md.",
    ]
    note_row = adj_row + 3
    for i, text in enumerate(notes):
        cell = ws.cell(row=note_row + i, column=1, value=text)
        cell.font = Font(bold=True) if i == 0 else Font(color="595959")
        cell.alignment = Alignment(wrap_text=False)


def main() -> None:
    with open(CSV_PATH, newline="", encoding="utf-8") as f:
        rows = list(csv.DictReader(f))

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    data = wb.create_sheet("Data")
    build_data_sheet(data, rows)
    build_summary_sheet(summary)
    wb.save(XLSX_PATH)
    print(f"Wrote {XLSX_PATH} ({len(rows)} data rows)")


if __name__ == "__main__":
    main()
